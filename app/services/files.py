"""File type detection and normalisation (Guide §10). Store large, reason small.

- Type is detected from magic bytes, never from the filename or the browser's MIME.
- Images: EXIF orientation fixed, converted to JPEG; a smaller copy goes to the AI.
- PDF: text extracted locally; for long PDFs only the most relevant pages are rendered for the AI
  and (if the original is large) packed into a compact derivative for Razorpay.
- DOCX: ZIP-bomb and macro checks, paragraphs + tables extracted, a PDF derivative generated.
Razorpay Documents accepts only JPG/JPEG/PNG/PDF, so everything else is converted.
"""
import io
import re
import zipfile
from dataclasses import dataclass, field

from app.errors import FILE_PROCESSING_FAILED, UNSUPPORTED_FILE_TYPE, ProofBridgeError

JPEG, PNG, PDF = "image/jpeg", "image/png", "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
HEIC = "image/heic"
ACCEPTED_DECLARED = {JPEG, PNG, PDF, DOCX, "image/jpg", HEIC, "image/heif"}

AI_IMAGE_MAX_PX = 1600
RZP_IMAGE_MAX_PX = 2400
RZP_PDF_MAX_BYTES = 8 * 1024 * 1024        # above this, send a compact relevant-pages derivative
PDF_AI_PAGES = 4
DOCX_MAX_EXPANDED = 60 * 1024 * 1024
DOCX_MAX_RATIO = 120

KEYWORDS = re.compile(r"(order|invoice|amount|₹|rs\.?|inr|awb|tracking|docket|deliver|courier|refund|"
                      r"customer|paid|payment|receipt|shipped|dispatch)", re.I)


def detect_mime(head: bytes) -> str | None:
    if head.startswith(b"\xff\xd8\xff"):
        return JPEG
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return PNG
    if head.startswith(b"%PDF-"):
        return PDF
    if head[4:12] in (b"ftypheic", b"ftypheix", b"ftypmif1", b"ftyphevc"):
        return HEIC
    if head.startswith(b"PK\x03\x04"):
        return "application/zip"     # refined to DOCX by inspect_docx
    return None


@dataclass
class Normalized:
    detected_mime: str
    ai_images: list[tuple[bytes, str]] = field(default_factory=list)
    ai_text: str | None = None
    rzp_bytes: bytes | None = None             # what goes to Razorpay Documents
    rzp_mime: str | None = None
    rzp_filename: str = "evidence"
    notes: list[str] = field(default_factory=list)


def _image(data: bytes, max_px: int, quality: int = 85) -> bytes:
    from PIL import Image, ImageOps
    with Image.open(io.BytesIO(data)) as im:
        im = ImageOps.exif_transpose(im)
        if im.mode not in ("RGB", "L"):
            im = im.convert("RGB")
        im.thumbnail((max_px, max_px))
        out = io.BytesIO()
        im.save(out, "JPEG", quality=quality, optimize=True)   # also strips EXIF (location etc.)
        return out.getvalue()


def normalize_image(data: bytes, base: str) -> Normalized:
    try:
        rzp = _image(data, RZP_IMAGE_MAX_PX)
        ai = _image(data, AI_IMAGE_MAX_PX, 80)
    except Exception:
        raise ProofBridgeError(FILE_PROCESSING_FAILED, "This image couldn't be opened. Try another photo.", 422)
    return Normalized(detected_mime=JPEG, ai_images=[(ai, JPEG)], rzp_bytes=rzp, rzp_mime=JPEG,
                      rzp_filename=f"{base}.jpg")


def normalize_pdf(data: bytes, base: str) -> Normalized:
    import fitz  # PyMuPDF
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception:
        raise ProofBridgeError(FILE_PROCESSING_FAILED, "This PDF couldn't be opened.", 422)
    if doc.needs_pass:
        raise ProofBridgeError(FILE_PROCESSING_FAILED, "This PDF is password-protected. Remove the password and retry.", 422)
    texts = [p.get_text("text") or "" for p in doc]
    if len(doc) <= PDF_AI_PAGES:
        pages = list(range(len(doc)))
    else:   # rank pages by dispute-relevant terms; keep reading order
        scored = sorted(range(len(doc)), key=lambda i: -len(KEYWORDS.findall(texts[i])))
        pages = sorted(scored[:PDF_AI_PAGES])
    images = []
    for i in pages:
        pix = doc[i].get_pixmap(dpi=110)
        images.append((_image(pix.tobytes("png"), AI_IMAGE_MAX_PX, 80), JPEG))
    text = "\n\n".join(f"[page {i + 1}]\n{texts[i].strip()}" for i in pages if texts[i].strip()) or None
    n = Normalized(detected_mime=PDF, ai_images=images, ai_text=text, rzp_mime=PDF, rzp_filename=f"{base}.pdf")
    if len(data) <= RZP_PDF_MAX_BYTES:
        n.rzp_bytes = data
    else:
        out = fitz.open()
        for i in pages:
            out.insert_pdf(doc, from_page=i, to_page=i)
        n.rzp_bytes = out.tobytes(garbage=3, deflate=True)
        n.notes.append(f"Large PDF: pages {', '.join(str(i + 1) for i in pages)} of {len(doc)} sent to Razorpay")
    return n


def inspect_docx(data: bytes) -> None:
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise ProofBridgeError(UNSUPPORTED_FILE_TYPE, "This file isn't a supported document.", 415)
    names = set(z.namelist())
    if "word/document.xml" not in names:
        raise ProofBridgeError(UNSUPPORTED_FILE_TYPE, "Only JPG, PNG, PDF and Word (.docx) files are supported.", 415)
    if any(n.lower().endswith("vbaproject.bin") for n in names):
        raise ProofBridgeError(UNSUPPORTED_FILE_TYPE, "Macro-enabled Word files aren't supported.", 415)
    if any(i.flag_bits & 0x1 for i in z.infolist()):
        raise ProofBridgeError(UNSUPPORTED_FILE_TYPE, "Encrypted Word files aren't supported.", 415)
    expanded = sum(i.file_size for i in z.infolist())
    if expanded > DOCX_MAX_EXPANDED or expanded > max(len(data), 1) * DOCX_MAX_RATIO:
        raise ProofBridgeError(UNSUPPORTED_FILE_TYPE, "This Word file is too large to process safely.", 415)


def _docx_to_pdf(paragraphs: list[str], tables: list[list[list[str]]], title: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape
    styles = getSampleStyleSheet()
    out = io.BytesIO()
    flow = [Paragraph(escape(title), styles["Heading2"]), Spacer(1, 8)]
    for p in paragraphs:
        flow.append(Paragraph(escape(p), styles["BodyText"]))
    for t in tables:
        if t:
            tbl = Table([[Paragraph(escape(c), styles["BodyText"]) for c in row] for row in t], repeatRows=1)
            tbl.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey)]))
            flow += [Spacer(1, 8), tbl]
    SimpleDocTemplate(out, pagesize=A4, title=title).build(flow)
    return out.getvalue()


def normalize_docx(data: bytes, base: str, original_name: str) -> Normalized:
    inspect_docx(data)
    try:
        import docx
        d = docx.Document(io.BytesIO(data))
        paras = [p.text.strip() for p in d.paragraphs if p.text.strip()]
        tables = [[[c.text.strip() for c in row.cells] for row in t.rows] for t in d.tables]
    except Exception:
        raise ProofBridgeError(FILE_PROCESSING_FAILED, "This Word file couldn't be read.", 422)
    table_text = "\n".join(" | ".join(r) for t in tables for r in t)
    text = "\n".join(paras) + ("\n" + table_text if table_text else "")
    pdf = _docx_to_pdf(paras, tables, f"Converted from {original_name}")
    return Normalized(detected_mime=DOCX, ai_text=text[:20000] or None, rzp_bytes=pdf, rzp_mime=PDF,
                      rzp_filename=f"{base}.pdf", notes=["Word file converted to PDF for Razorpay"])


def normalize(data: bytes, original_name: str) -> Normalized:
    mime = detect_mime(data[:16])
    base = re.sub(r"[^A-Za-z0-9_-]+", "_", original_name.rsplit(".", 1)[0])[:60] or "evidence"
    if mime in (JPEG, PNG):
        return normalize_image(data, base)
    if mime == PDF:
        return normalize_pdf(data, base)
    if mime == "application/zip":
        return normalize_docx(data, base, original_name)
    if mime == HEIC:
        raise ProofBridgeError(UNSUPPORTED_FILE_TYPE, "iPhone HEIC photos aren't supported yet. Share it as JPG "
                               "(Settings > Camera > Formats > Most Compatible) or take a screenshot.", 415)
    raise ProofBridgeError(UNSUPPORTED_FILE_TYPE, "Only JPG, PNG, PDF and Word (.docx) files are supported.", 415)
