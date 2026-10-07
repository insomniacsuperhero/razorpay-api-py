"""Synthetic evidence files generated in memory (no real customer data)."""
import io


def jpeg(color=(240, 240, 240), size=(900, 1600), text_lines=()) -> bytes:
    from PIL import Image, ImageDraw
    im = Image.new("RGB", size, color)
    d = ImageDraw.Draw(im)
    for i, line in enumerate(text_lines):
        d.text((40, 60 + i * 40), line, fill=(20, 20, 20))
    out = io.BytesIO()
    im.save(out, "JPEG")
    return out.getvalue()


def png() -> bytes:
    from PIL import Image
    out = io.BytesIO()
    Image.new("RGBA", (400, 300), (0, 128, 255, 255)).save(out, "PNG")
    return out.getvalue()


def pdf(lines: list[str], pages: int = 1, filler: str = "Lorem ipsum terms page") -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    out = io.BytesIO()
    c = canvas.Canvas(out, pagesize=A4)
    for p in range(pages):
        if p == pages - 2 or pages == 1:
            for i, line in enumerate(lines):
                c.drawString(60, 780 - i * 20, line)
        else:
            c.drawString(60, 780, f"{filler} {p + 1}")
        c.showPage()
    c.save()
    return out.getvalue()


def docx(lines: list[str]) -> bytes:
    import docx as d
    doc = d.Document()
    for line in lines:
        doc.add_paragraph(line)
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "From", "Message"
    t.cell(1, 0).text, t.cell(1, 1).text = "Customer", "Received, thank you"
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
